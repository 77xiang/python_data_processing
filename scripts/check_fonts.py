# -*- coding: utf-8 -*-
"""检查报告 docx 中每一处文字的字体是否符合规范。

规范：
  - 正文 / 标题 / 表格数据：西文(ascii/hAnsi/cs) = Times New Roman，中文(eastAsia) = 宋体 或 黑体
  - 代码块：Courier New（刻意例外）
  - 不允许出现：西文字体缺失（会回退到中文字体，导致数字/英文不是 TNR）或其他西文字体
"""
import sys
from collections import Counter

import docx
from docx.oxml.ns import qn

PATH = "deliverables/城市电力负荷分析与预测_项目报告.docx"
TNR = "Times New Roman"
CODE = "Courier New"
CJK_EA = {"宋体", "黑体", "SimSun", "SimHei"}

CJK = re = None
import re  # noqa: E402

CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\u3000-\u303f\uff00-\uffef]")


def iter_paragraphs(doc):
    for para in doc.paragraphs:
        yield ("body", para)
    for ti, t in enumerate(doc.tables):
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield ("table%d" % ti, para)
    for s in doc.sections:
        for para in s.header.paragraphs:
            yield ("header", para)
        for para in s.footer.paragraphs:
            yield ("footer", para)


NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def load_theme(path):
    """读取主题字体：Word 保存时会把显式字体改写成 asciiTheme=minorHAnsi 这类引用。"""
    import zipfile
    from lxml import etree
    m = {}
    with zipfile.ZipFile(path) as z:
        if "word/theme/theme1.xml" in z.namelist():
            root = etree.fromstring(z.read("word/theme/theme1.xml"))
            for tag, prefix in (("minorFont", "minor"), ("majorFont", "major")):
                for f in root.iter(NS_A + tag):
                    for c in f:
                        ln = etree.QName(c).localname
                        if ln == "latin":
                            m[prefix + "HAnsi"] = c.get("typeface")
                            m[prefix + "Ascii"] = c.get("typeface")
                        elif ln == "ea":
                            m[prefix + "EastAsia"] = c.get("typeface")
                        elif ln == "cs":
                            m[prefix + "Bidi"] = c.get("typeface")
    return m


def fonts_of(run, theme=None):
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None, None
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        return None, None
    theme = theme or {}

    def res(attr, theme_attr):
        v = rf.get(qn("w:" + attr))
        if v:
            return v
        tv = rf.get(qn("w:" + theme_attr))
        if tv:
            return theme.get(tv)
        return None

    ascii_f = res("ascii", "asciiTheme")
    ea_f = res("eastAsia", "eastAsiaTheme")
    return (ascii_f, ea_f, res("hAnsi", "hAnsiTheme"), res("cs", "cstheme")), (ascii_f, ea_f)


def main():
    d = docx.Document(PATH)
    styles = d.styles
    theme = load_theme(PATH)
    print("主题字体映射:", theme)
    # 1) 文档级默认字体（docDefaults 位于 styles 部件，不在 document.xml）
    dd = d.styles.element.find(qn("w:docDefaults"))
    print("=== docDefaults ===")
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault"))
        rf = rpr.find(qn("w:rPr")).find(qn("w:rFonts")) if rpr is not None and rpr.find(qn("w:rPr")) is not None else None
        if rf is not None:
            print("  ascii=%s  eastAsia=%s  hAnsi=%s  cs=%s" % (
                rf.get(qn("w:ascii")), rf.get(qn("w:eastAsia")),
                rf.get(qn("w:hAnsi")), rf.get(qn("w:cs"))))
        else:
            print("  !! 无 rFonts，西文将回退到 Word 默认")
    for nm in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Body Text", "HTML Code"):
        try:
            st = styles[nm]
        except KeyError:
            print("  样式缺失:", nm)
            continue
        rf = st.element.find(qn("w:rPr"))
        rf = rf.find(qn("w:rFonts")) if rf is not None else None
        if rf is not None:
            print("  %-10s ascii=%s eastAsia=%s" % (nm, rf.get(qn("w:ascii")), rf.get(qn("w:eastAsia"))))

    print("\n=== run 级检查 ===")
    stat = Counter()
    problems = []
    n_run = 0
    n_digit_en_run = 0
    for scope, para in iter_paragraphs(d):
        stname = para.style.name if para.style else ""
        for run in para.runs:
            txt = run.text
            if not txt or not txt.strip():
                continue
            n_run += 1
            f4, f2 = fonts_of(run, theme)
            ascii_f, ea_f = f2 if f2 else (None, None)
            is_code = (ascii_f == CODE) or ("Code" in stname)
            has_ascii = bool(re.search(r"[0-9A-Za-z]", txt))
            has_cjk = bool(CJK_RE.search(txt))
            if has_ascii:
                n_digit_en_run += 1
            if is_code:
                stat["code"] += 1
                if ascii_f != CODE:
                    problems.append((scope, "代码段西文非 Courier New: %s" % ascii_f, txt[:40]))
                continue
            if ascii_f is None:
                stat["ascii缺失(继承)"] += 1
                problems.append((scope, "西文字体未指定(rFonts/ascii 缺失)", txt[:40]))
            elif ascii_f != TNR:
                stat["西文非TNR:" + ascii_f] += 1
                problems.append((scope, "西文字体=%s" % ascii_f, txt[:40]))
            else:
                stat["西文TNR-OK"] += 1
            if has_cjk:
                if ea_f is None:
                    stat["eastAsia缺失"] += 1
                    problems.append((scope, "中文字体未指定", txt[:40]))
                elif ea_f not in CJK_EA:
                    stat["中文非宋黑:" + str(ea_f)] += 1
                    problems.append((scope, "中文字体=%s" % ea_f, txt[:40]))
                else:
                    stat["中文OK"] += 1

    print("含文字的 run 总数: %d（其中含数字/英文的 run: %d）" % (n_run, n_digit_en_run))
    for k, v in stat.most_common():
        print("  %-24s %d" % (k, v))
    print("\n问题数: %d" % len(problems))
    for scope, why, t in problems[:40]:
        print("  [%s] %s  |  %s" % (scope, why, t))
    if len(problems) > 40:
        print("  ... 其余 %d 条省略" % (len(problems) - 40))
    return 0 if not problems else 1


if __name__ == "__main__":
    sys.exit(main())
